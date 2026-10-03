"""
仓库管理视图
"""
import logging
import io
from datetime import date

from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    TransferUnit, TransferUnitNameRecord, Authorization, ReturnApproval, Handover,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    TransferUnitSerializer, TransferUnitCreateSerializer, UnitRenameSerializer,
    UnitStatusSerializer, AuthorizationSerializer, AuthorizationWriteSerializer,
    AuthorizationRevokeSerializer, ReturnApprovalSerializer, ReturnApprovalWriteSerializer,
    AdmissionEvaluateSerializer, HandoverReceiveSerializer, HandoverSerializer,
)
from .admission import evaluate_admission, record_handover

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 移交单位准入管理 ====================

def _first_error(serializer):
    """提取序列化器第一条错误信息"""
    errors = serializer.errors
    value = list(errors.values())[0]
    if isinstance(value, list):
        value = value[0]
    return str(value)


def _paginate(request, queryset):
    page = max(int(request.query_params.get('page', 1)), 1)
    page_size = max(int(request.query_params.get('page_size', 10)), 1)
    start = (page - 1) * page_size
    return queryset.count(), queryset[start:start + page_size], page, page_size


class TransferUnitListView(APIView):
    """移交单位列表/创建"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = TransferUnit.objects.prefetch_related('name_records').all()
        name = request.query_params.get('name')
        if name:
            queryset = queryset.filter(name__icontains=name)
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        queryset = queryset.order_by('-created_at')

        total, units, page, page_size = _paginate(request, queryset)
        return success_response(data={
            'list': TransferUnitSerializer(units, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })

    def post(self, request):
        serializer = TransferUnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        with transaction.atomic():
            unit = TransferUnit.objects.create(
                code=data['code'], name=data['name'], created_by=request.user,
            )
            # 当前名称同时登记为首个名称版本，保证名称版本链完整
            TransferUnitNameRecord.objects.create(
                unit=unit, name=unit.name, effective_from=timezone.localdate(),
                reason='单位建档', created_by=request.user,
            )

        logger.info(f"User {request.user.username} created transfer unit {unit.name}")
        return success_response(data=TransferUnitSerializer(unit).data, message='创建成功')


class TransferUnitDetailView(APIView):
    """移交单位详情"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        unit = TransferUnit.objects.prefetch_related('name_records').filter(pk=pk).first()
        if unit is None:
            return error_response(message='移交单位不存在', code=404)
        return success_response(data=TransferUnitSerializer(unit).data)


class TransferUnitRenameView(APIView):
    """移交单位名称变更：旧名称保留为历史名称，新名称不得与任何在册/历史名称冲突"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        unit = TransferUnit.objects.filter(pk=pk).first()
        if unit is None:
            return error_response(message='移交单位不存在', code=404)

        serializer = UnitRenameSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        with transaction.atomic():
            old_name = unit.name
            unit.name = data['new_name']
            unit.save(update_fields=['name', 'updated_at'])
            TransferUnitNameRecord.objects.create(
                unit=unit, name=data['new_name'], effective_from=data['effective_from'],
                reason=data.get('reason', ''), created_by=request.user,
            )

        logger.info(
            f"User {request.user.username} renamed transfer unit {old_name} -> {unit.name}"
        )
        return success_response(data=TransferUnitSerializer(unit).data, message='名称变更成功')


class TransferUnitStatusView(APIView):
    """移交单位暂停/恢复/终止"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        unit = TransferUnit.objects.filter(pk=pk).first()
        if unit is None:
            return error_response(message='移交单位不存在', code=404)

        serializer = UnitStatusSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        unit.status = serializer.validated_data['status']
        unit.suspend_reason = serializer.validated_data.get('reason', '')
        unit.suspended_at = timezone.now() if unit.is_suspended else None
        unit.save(update_fields=['status', 'suspend_reason', 'suspended_at', 'updated_at'])

        logger.info(
            f"User {request.user.username} set transfer unit {unit.name} status {unit.status}"
        )
        return success_response(data=TransferUnitSerializer(unit).data, message='状态已更新')


class AuthorizationListView(APIView):
    """准入授权列表/创建"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Authorization.objects.prefetch_related('categories').all()
        unit_id = request.query_params.get('unit')
        if unit_id:
            queryset = queryset.filter(unit_id=unit_id)
        kind = request.query_params.get('kind')
        if kind:
            queryset = queryset.filter(kind=kind)
        queryset = queryset.order_by('-valid_from', '-id')

        total, authorizations, page, page_size = _paginate(request, queryset)
        return success_response(data={
            'list': AuthorizationSerializer(authorizations, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })

    def post(self, request):
        serializer = AuthorizationWriteSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        with transaction.atomic():
            authorization = Authorization.objects.create(
                unit_id=data['unit'], kind=data['kind'],
                document_no=data.get('document_no', ''),
                valid_from=data['valid_from'], valid_to=data['valid_to'],
                remark=data.get('remark', ''), created_by=request.user,
            )
            authorization.categories.set(data['category_ids'])

        logger.info(
            f"User {request.user.username} created {authorization.kind} authorization "
            f"for unit {authorization.unit_id}"
        )
        return success_response(
            data=AuthorizationSerializer(authorization).data, message='授权登记成功'
        )


class AuthorizationDetailView(APIView):
    """准入授权详情"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        authorization = Authorization.objects.prefetch_related('categories').filter(pk=pk).first()
        if authorization is None:
            return error_response(message='准入授权不存在', code=404)
        return success_response(data=AuthorizationSerializer(authorization).data)


class AuthorizationRevokeView(APIView):
    """撤销准入授权（已用于接收记录的授权仍保留，仅停止后续适用）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        authorization = Authorization.objects.filter(pk=pk).first()
        if authorization is None:
            return error_response(message='准入授权不存在', code=404)
        if authorization.is_revoked:
            return error_response(message='该授权已撤销，无需重复操作')

        serializer = AuthorizationRevokeSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        authorization.revoked_at = timezone.now()
        authorization.revoke_reason = serializer.validated_data.get('reason', '')
        authorization.save(update_fields=['revoked_at', 'revoke_reason', 'updated_at'])
        logger.info(f"User {request.user.username} revoked authorization {pk}")
        return success_response(data=AuthorizationSerializer(authorization).data, message='授权已撤销')


class ReturnApprovalListView(APIView):
    """退回批准单列表/创建"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = ReturnApproval.objects.select_related('unit', 'category').all()
        unit_id = request.query_params.get('unit')
        if unit_id:
            queryset = queryset.filter(unit_id=unit_id)
        status_value = request.query_params.get('status')
        if status_value:
            queryset = queryset.filter(status=status_value)
        queryset = queryset.order_by('-approved_at')

        total, approvals, page, page_size = _paginate(request, queryset)
        return success_response(data={
            'list': ReturnApprovalSerializer(approvals, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })

    def post(self, request):
        serializer = ReturnApprovalWriteSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        approval = ReturnApproval.objects.create(
            unit_id=data['unit'], document_no=data['document_no'],
            category_id=data.get('category'), quantity=data.get('quantity'),
            reason=data.get('reason', ''), approved_by=request.user,
        )
        logger.info(
            f"User {request.user.username} created return approval {approval.document_no}"
        )
        return success_response(
            data=ReturnApprovalSerializer(approval).data, message='退回批准单登记成功'
        )


def _decision_payload(decision, category):
    """构造判定结果响应体"""
    return {
        'accepted': decision.accepted,
        'reason_code': decision.reason_code,
        'message': decision.reason_message,
        'purpose': decision.purpose,
        'on_date': decision.on_date.isoformat(),
        'unit_name': decision.unit_name,
        'unit_id': decision.unit.id if decision.unit else None,
        'unit_status': decision.unit.get_status_display() if decision.unit else '',
        'name_match_type': decision.name_match_type,
        'matched_name': (
            decision.matched_name_record.name if decision.matched_name_record else ''
        ),
        'category_id': category.id,
        'category_name': category.name,
        'authorization_id': decision.authorization.id if decision.authorization else None,
        'authorization_kind': (
            decision.authorization.get_kind_display() if decision.authorization else ''
        ),
        'authorization_document_no': (
            decision.authorization.document_no if decision.authorization else ''
        ),
        'return_approval_id': decision.return_approval.id if decision.return_approval else None,
        'return_document_no': (
            decision.return_approval.document_no if decision.return_approval else ''
        ),
        'basis': decision.basis,
    }


class AdmissionEvaluateView(APIView):
    """收件准入预检：只判定，不写台账，返回通过/拒绝的具体依据"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = AdmissionEvaluateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        category = Category.objects.get(pk=data['category'])
        decision = evaluate_admission(
            unit_name=data['unit_name'], category=category, on_date=data['on_date'],
            purpose=data.get('purpose', 'handover'),
            return_document_no=data.get('return_document_no'),
        )
        return success_response(data=_decision_payload(decision, category))


class HandoverReceiveView(APIView):
    """正式接收：执行准入判定，通过才写入正式台账；拒绝同样留痕并返回依据"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = HandoverReceiveSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))

        data = serializer.validated_data
        category = Category.objects.get(pk=data['category'])
        variety = None
        if data.get('variety'):
            variety = Variety.objects.filter(pk=data['variety']).first()

        decision = evaluate_admission(
            unit_name=data['unit_name'], category=category, on_date=data['on_date'],
            purpose=data.get('purpose', 'handover'),
            return_document_no=data.get('return_document_no'),
        )
        handover = record_handover(
            decision=decision, category=category, variety=variety,
            material_name=data.get('material_name', ''),
            material_code=data.get('material_code', ''),
            quantity=data.get('quantity', 0), batch_no=data.get('batch_no', ''),
            operator=request.user,
        )
        payload = _decision_payload(decision, category)
        payload['handover'] = HandoverSerializer(handover).data

        if not decision.accepted:
            logger.info(
                f"User {request.user.username} rejected handover from "
                f"{decision.unit_name}: {decision.reason_code}"
            )
            return error_response(
                message=f'接收拒绝：{decision.reason_message}', code=200, data=payload
            )

        logger.info(
            f"User {request.user.username} accepted {decision.purpose} from "
            f"{decision.unit.name} via {decision.authorization or decision.return_approval}"
        )
        return success_response(data=payload, message='接收通过，已登记正式台账')


class HandoverListView(APIView):
    """移交接收记录查询（台账）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Handover.objects.select_related(
            'unit', 'category', 'variety', 'authorization', 'return_approval', 'received_by'
        ).all()

        unit_name = request.query_params.get('unit_name')
        if unit_name:
            queryset = queryset.filter(
                Q(unit__name__icontains=unit_name)
                | Q(unit_name_snapshot__icontains=unit_name)
            )
        decision_value = request.query_params.get('decision')
        if decision_value:
            queryset = queryset.filter(decision=decision_value)
        purpose_value = request.query_params.get('purpose')
        if purpose_value:
            queryset = queryset.filter(purpose=purpose_value)
        category_id = request.query_params.get('category')
        if category_id:
            queryset = queryset.filter(category_id=category_id)
        start_date = request.query_params.get('start_date')
        if start_date:
            queryset = queryset.filter(handover_date__gte=start_date)
        end_date = request.query_params.get('end_date')
        if end_date:
            queryset = queryset.filter(handover_date__lte=end_date)
        queryset = queryset.order_by('-created_at')

        total, handovers, page, page_size = _paginate(request, queryset)
        return success_response(data={
            'list': HandoverSerializer(handovers, many=True).data,
            'total': total, 'page': page, 'page_size': page_size,
        })


class HandoverDetailView(APIView):
    """单次接收记录详情：含逐条判定依据"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        handover = Handover.objects.select_related(
            'unit', 'category', 'authorization', 'return_approval', 'received_by'
        ).filter(pk=pk).first()
        if handover is None:
            return error_response(message='接收记录不存在', code=404)
        return success_response(data=HandoverSerializer(handover).data)
